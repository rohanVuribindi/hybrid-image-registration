"""
Convert ZENITH markdown documentation to DOCX and PDF formats.
Outputs:
- docs/ZENITH_COMPLETE_TECHNICAL_DOCUMENTATION.docx
- docs/ZENITH_COMPLETE_TECHNICAL_DOCUMENTATION.pdf
"""

import os
import sys
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
)
from reportlab.pdfgen import canvas

def clean_xml_string(s):
    if not isinstance(s, str):
        return ""
    # Remove control characters that are invalid in XML (keep tab \t, newline \n, cr \r)
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', s)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def convert_md_to_docx(md_path, docx_path):
    print(f"Converting {md_path} to DOCX: {docx_path}...")
    doc = Document()
    
    # Page setup (margins 0.75 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "ZENITH — Planetary Cross-Modal Image Registration | Technical Specification"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.runs[0].font.size = Pt(8.5)
        hp.runs[0].font.color.rgb = RGBColor(120, 130, 150)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "CONFIDENTIAL & PROPRIETARY — ZENITH MISSION SYSTEM v2.0"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.runs[0].font.size = Pt(8.5)
        fp.runs[0].font.color.rgb = RGBColor(140, 140, 140)

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        rows = []
        for line in tbl_lines:
            line_str = line.strip()
            if line_str.startswith("|") and line_str.endswith("|"):
                parts = [clean_xml_string(p.strip()) for p in line_str.strip("|").split("|")]
                if all(re.match(r'^:?-+:?$', p) for p in parts):
                    continue
                rows.append(parts)
        if not rows:
            return
        
        col_count = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=col_count)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = 'Table Grid'
        
        for r_idx, r_data in enumerate(rows):
            for c_idx, cell_value in enumerate(r_data):
                if c_idx < col_count:
                    cell = t.cell(r_idx, c_idx)
                    cell.text = cell_value
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    p.runs[0].font.size = Pt(9.5)
                    if r_idx == 0:
                        p.runs[0].font.bold = True
                        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                        set_cell_background(cell, "1A2B4C")
                    else:
                        if r_idx % 2 == 1:
                            set_cell_background(cell, "F4F6F9")
                        else:
                            set_cell_background(cell, "FFFFFF")
        doc.add_paragraph()

    for line in lines:
        stripped = clean_xml_string(line.strip())
        
        # Code block handling
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                code_text = clean_xml_string("\n".join(code_lines))
                code_p = doc.add_paragraph()
                code_p.paragraph_format.left_indent = Inches(0.25)
                code_p.paragraph_format.space_before = Pt(4)
                code_p.paragraph_format.space_after = Pt(4)
                run = code_p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(30, 40, 60)
                code_lines = []
            else:
                if in_table:
                    flush_table(table_lines)
                    table_lines = []
                    in_table = False
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line.rstrip("\n"))
            continue

        # Table handling
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            flush_table(table_lines)
            table_lines = []
            in_table = False

        if not stripped:
            continue

        # Headings
        if stripped.startswith("# "):
            h = doc.add_heading(level=0)
            run = h.add_run(stripped[2:])
            run.font.name = 'Arial'
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(14, 30, 64)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(6)
        elif stripped.startswith("## "):
            h = doc.add_heading(level=1)
            run = h.add_run(stripped[3:])
            run.font.name = 'Arial'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(24, 52, 98)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
        elif stripped.startswith("### "):
            h = doc.add_heading(level=2)
            run = h.add_run(stripped[4:])
            run.font.name = 'Arial'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(40, 70, 120)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
        elif stripped.startswith("#### "):
            h = doc.add_heading(level=3)
            run = h.add_run(stripped[5:])
            run.font.name = 'Arial'
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(60, 80, 110)
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[2:])
            run.font.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(80, 60, 20)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(stripped[2:])
            run.font.size = Pt(10)
        elif re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            text_part = re.sub(r'^\d+\.\s*', '', stripped)
            run = p.add_run(text_part)
            run.font.size = Pt(10)
        elif stripped == "---":
            continue
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(stripped)
            run.font.name = 'Calibri'
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(30, 30, 30)

    if in_table:
        flush_table(table_lines)

    doc.save(docx_path)
    print(f"Saved DOCX: {docx_path}")

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_header_footer(num_pages)
            super().showPage()
        super().save()

    def draw_header_footer(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#667788"))
        
        # Header (pages > 1)
        if self._pageNumber > 1:
            self.drawString(54, 750, "ZENITH: Autonomous Planetary Cross-Modal Image Registration")
            self.drawRightString(612 - 54, 750, "Technical Specification v2.0")
            self.setStrokeColor(colors.HexColor("#CCD5E0"))
            self.setLineWidth(0.5)
            self.line(54, 744, 612 - 54, 744)

        # Footer (all pages)
        self.setStrokeColor(colors.HexColor("#CCD5E0"))
        self.setLineWidth(0.5)
        self.line(54, 45, 612 - 54, 45)
        self.drawString(54, 32, "ZENITH Technical Documentation — CONFIDENTIAL & PROPRIETARY")
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(612 - 54, 32, page_text)
        self.restoreState()

def convert_md_to_pdf(md_path, pdf_path):
    print(f"Converting {md_path} to PDF: {pdf_path}...")
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=60,
        bottomMargin=60
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0D2240"),
        spaceAfter=6
    )
    
    h1_style = ParagraphStyle(
        'SectionHeading1',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#1A365D"),
        spaceBefore=12,
        spaceAfter=5,
        keepWithNext=True
    )
    
    h2_style = ParagraphStyle(
        'SectionHeading2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#2B6CB0"),
        spaceBefore=8,
        spaceAfter=3,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#2D3748"),
        spaceAfter=3
    )
    
    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=body_style,
        leftIndent=12,
        firstLineIndent=-8,
        spaceAfter=2
    )

    code_style = ParagraphStyle(
        'CodeBlock',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor("#1A202C"),
        backColor=colors.HexColor("#F7FAFC"),
        borderColor=colors.HexColor("#E2E8F0"),
        borderWidth=0.5,
        borderPadding=4,
        spaceBefore=4,
        spaceAfter=5
    )
    
    table_cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=9.5,
        textColor=colors.HexColor("#2D3748")
    )
    
    table_hdr_style = ParagraphStyle(
        'TableHdr',
        parent=table_cell_style,
        fontName='Helvetica-Bold',
        textColor=colors.white
    )

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    story = []
    in_code = False
    code_buf = []
    in_tbl = False
    tbl_buf = []

    def flush_pdf_table(tbl_lines):
        if not tbl_lines:
            return
        rows = []
        for line in tbl_lines:
            line_str = line.strip()
            if line_str.startswith("|") and line_str.endswith("|"):
                parts = [clean_xml_string(p.strip()) for p in line_str.strip("|").split("|")]
                if all(re.match(r'^:?-+:?$', p) for p in parts):
                    continue
                rows.append(parts)
        if not rows:
            return

        col_count = max(len(r) for r in rows)
        avail_width = 504
        col_width = avail_width / col_count

        table_data = []
        for r_idx, r in enumerate(rows):
            row_cells = []
            for c_idx in range(col_count):
                val = r[c_idx] if c_idx < len(r) else ""
                clean_val = val.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                st = table_hdr_style if r_idx == 0 else table_cell_style
                row_cells.append(Paragraph(clean_val, st))
            table_data.append(row_cells)

        t = Table(table_data, colWidths=[col_width]*col_count)
        t_style = [
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1A365D")),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('LEFTPADDING', (0,0), (-1,-1), 3),
            ('RIGHTPADDING', (0,0), (-1,-1), 3),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E0")),
        ]
        for r_i in range(1, len(table_data)):
            if r_i % 2 == 1:
                t_style.append(('BACKGROUND', (0, r_i), (-1, r_i), colors.HexColor("#F7FAFC")))
        t.setStyle(TableStyle(t_style))
        story.append(Spacer(1, 3))
        story.append(t)
        story.append(Spacer(1, 4))

    for line in lines:
        stripped = clean_xml_string(line.strip())
        
        if stripped.startswith("```"):
            if in_code:
                in_code = False
                code_text = "<br/>".join(
                    c.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace(" ", "&nbsp;")
                    for c in code_buf
                )
                story.append(Paragraph(code_text, code_style))
                code_buf = []
            else:
                if in_tbl:
                    flush_pdf_table(tbl_buf)
                    tbl_buf = []
                    in_tbl = False
                in_code = True
                code_buf = []
            continue

        if in_code:
            code_buf.append(line.rstrip("\n"))
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            in_tbl = True
            tbl_buf.append(stripped)
            continue
        elif in_tbl:
            flush_pdf_table(tbl_buf)
            tbl_buf = []
            in_tbl = False

        if not stripped:
            continue

        clean_text = stripped.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        clean_text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', clean_text)
        clean_text = re.sub(r'\*(.*?)\*', r'<i>\1</i>', clean_text)
        clean_text = re.sub(r'`(.*?)`', r'<font face="Courier" color="#1A202C">\1</font>', clean_text)

        if stripped.startswith("# "):
            story.append(Paragraph(clean_text[2:], title_style))
            story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1A365D"), spaceAfter=8))
        elif stripped.startswith("## "):
            story.append(Paragraph(clean_text[3:], h1_style))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#CBD5E0"), spaceAfter=4))
        elif stripped.startswith("### "):
            story.append(Paragraph(clean_text[4:], h2_style))
        elif stripped.startswith("- ") or stripped.startswith("* "):
            story.append(Paragraph(f"&bull; {clean_text[2:]}", bullet_style))
        elif re.match(r'^\d+\.\s', stripped):
            story.append(Paragraph(clean_text, bullet_style))
        elif stripped == "---":
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#E2E8F0"), spaceAfter=6, spaceBefore=6))
        else:
            story.append(Paragraph(clean_text, body_style))

    if in_tbl:
        flush_pdf_table(tbl_buf)

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"Saved PDF: {pdf_path}")

if __name__ == "__main__":
    md_file = os.path.join("docs", "ZENITH_COMPLETE_TECHNICAL_DOCUMENTATION.md")
    docx_file = os.path.join("docs", "ZENITH_COMPLETE_TECHNICAL_DOCUMENTATION.docx")
    pdf_file = os.path.join("docs", "ZENITH_COMPLETE_TECHNICAL_DOCUMENTATION.pdf")
    
    if os.path.exists(md_file):
        convert_md_to_docx(md_file, docx_file)
        convert_md_to_pdf(md_file, pdf_file)
        print("Documentation generation complete!")
    else:
        print(f"Error: {md_file} not found.")
