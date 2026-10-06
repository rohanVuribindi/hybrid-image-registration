import pptx
from pptx.util import Inches, Pt

prs = pptx.Presentation(r'c:\Users\rohan\Downloads\LUNAR-MATCH_SIH2026_v5.pptx')
for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx+1} ====================")
    for s in slide.shapes:
        left_in = s.left / Inches(1) if s.left else 0
        top_in = s.top / Inches(1) if s.top else 0
        w_in = s.width / Inches(1) if s.width else 0
        h_in = s.height / Inches(1) if s.height else 0
        print(f"Shape ID {s.shape_id}: '{s.name}' | Type: {s.shape_type} | Box: ({left_in:.2f}, {top_in:.2f}, {w_in:.2f}, {h_in:.2f})")
        if s.has_text_frame:
            for p_idx, p in enumerate(s.text_frame.paragraphs):
                font_name = p.font.name if p.font else None
                font_size = p.font.size.pt if p.font and p.font.size else None
                bold = p.font.bold if p.font else None
                print(f"   P{p_idx} (lvl={p.level}, sz={font_size}, b={bold}): {p.text}")
