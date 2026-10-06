import pptx

prs = pptx.Presentation(r'c:\Users\rohan\Downloads\LUNAR-MATCH_SIH2026_v5.pptx')
with open(r'c:\Users\rohan\Downloads\prototype\slide_dump.txt', 'w', encoding='utf-8') as f:
    for idx, slide in enumerate(prs.slides):
        f.write(f'=== SLIDE {idx+1} ===\n')
        for shape in slide.shapes:
            f.write(f'Shape ID: {shape.shape_id}, Name: {shape.name}, Type: {shape.shape_type}\n')
            if shape.has_text_frame:
                for p_idx, p in enumerate(shape.text_frame.paragraphs):
                    runs_text = ''.join([r.text for r in p.runs])
                    f.write(f'   [P{p_idx}] text="{p.text}"\n')
            elif shape.has_table:
                f.write(f'   [TABLE] {len(shape.table.rows)}x{len(shape.table.columns)}\n')
                for r_idx, r in enumerate(shape.table.rows):
                    row_str = ' | '.join([c.text.strip().replace('\n', ' ') for c in r.cells])
                    f.write(f'     R{r_idx}: {row_str}\n')
print('Dumped successfully!')
